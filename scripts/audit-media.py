"""Inventory originals and make inspection sheets; never alters source pixels."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path('.media-tools').resolve()))
import hashlib, json, io
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation

root=Path('labani_di'); out=Path('.media-review');out.mkdir(exist_ok=True)
groups={}
for p in sorted(root.rglob('*')):
    if p.is_file(): groups.setdefault(hashlib.sha256(p.read_bytes()).hexdigest(),[]).append(p)
records=[]
for sha,paths in groups.items():
    p=paths[0]
    if p.suffix.lower() not in ['.jpg','.jpeg','.png','.webp']: continue
    im=ImageOps.exif_transpose(Image.open(p)); im.load()
    records.append(dict(id=f'img-{len(records)+1:04}',sha256=sha,paths=[str(q) for q in paths],width=im.width,height=im.height,format=im.format))
(out/'inventory.json').write_text(json.dumps(records,indent=2),encoding='utf-8')
for start in range(0,len(records),12):
    sheet=Image.new('RGB',(1600,1500),'#dddddd'); draw=ImageDraw.Draw(sheet)
    for j,r in enumerate(records[start:start+12]):
        im=ImageOps.exif_transpose(Image.open(r['paths'][0])).convert('RGB'); im.thumbnail((386,450))
        x=(j%4)*400;y=(j//4)*500
        sheet.paste(im,(x+(400-im.width)//2,y+(450-im.height)//2))
        draw.text((x+6,y+458),f"{r['id']}  {r['width']}x{r['height']}",fill='black')
        draw.text((x+6,y+475),Path(r['paths'][0]).name[:46],fill='black')
    sheet.save(out/f'sheet-{start//12+1:02}.jpg',quality=92)
pres=Presentation(next(root.glob('*.pptx'))); slides=[]
extract=out/'ppt-images';extract.mkdir(exist_ok=True)
for n,slide in enumerate(pres.slides,1):
    texts=[]; pictures=[]
    for shape in slide.shapes:
        if shape.has_text_frame: texts.append(shape.text)
        if shape.shape_type==13:
            dest=extract/f'slide-{n:02}-image-{len(pictures)+1:02}.{shape.image.ext}';dest.write_bytes(shape.image.blob)
            pictures.append(str(dest))
    slides.append(dict(slide=n,text='\n'.join(texts),images=pictures,notes=slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''))
(out/'portfolio.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
(out/'portfolio-text.txt').write_text('\n\n'.join(f"SLIDE {s['slide']}\n{s['text']}\nNOTES: {s['notes']}" for s in slides),encoding='utf-8')
print(f'{len(records)} distinct images; {(len(records)+11)//12} sheets; {len(slides)} slides')
